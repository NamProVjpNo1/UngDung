import hashlib
import io
import uuid
from pathlib import Path
import pymupdf
from docx import Document as DocxDocument
from fastapi import HTTPException
from sqlalchemy import select
from app.db.session import DATA_DIR
from app.models.entities import Chunk, Document
from app.services.exams import now

MAX_UPLOAD = 20 * 1024 * 1024


def extract_chunks(content, suffix):
    try:
        if suffix == ".pdf":
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.page_count > 1000:
                    raise ValueError("Tối đa 1.000 trang.")
                pages = [(i + 1, page.get_text()) for i, page in enumerate(pdf)]
        else:
            pages = [(1, "\n".join(p.text for p in DocxDocument(io.BytesIO(content)).paragraphs))]
        chunks = []
        for page, text in pages:
            text = text.strip()
            for offset in range(0, len(text), 2500):
                part = text[offset : offset + 2500].strip()
                if part:
                    chunks.append((page, part))
        if not chunks:
            raise ValueError("Không trích xuất được văn bản. PDF ảnh quét cần OCR riêng.")
        return chunks
    except Exception as exc:
        raise HTTPException(422, "Không đọc được tài liệu hoặc không có văn bản. PDF ảnh quét cần OCR riêng.") from exc


def import_document(db, filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix not in (".pdf", ".docx"):
        raise HTTPException(415, "Chỉ hỗ trợ PDF và DOCX.")
    if not content or len(content) > MAX_UPLOAD:
        raise HTTPException(413, "Tệp phải từ 1 byte đến 20 MB.")
    if suffix == ".docx":
        import zipfile

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024:
                    raise HTTPException(413, "Tài liệu giải nén vượt 80 MB.")
        except zipfile.BadZipFile as exc:
            raise HTTPException(422, "Tệp DOCX không hợp lệ.") from exc
    digest = hashlib.sha256(content).hexdigest()
    if db.scalar(select(Document).where(Document.ma_bam == digest)):
        raise HTTPException(409, "Tài liệu này đã được nhập (SHA-256 trùng nhau).")
    chunks = extract_chunks(content, suffix)
    relative = f"documents/{uuid.uuid4().hex}{suffix}"
    path = DATA_DIR / relative
    path.parent.mkdir(exist_ok=True)
    path.write_bytes(content)
    doc = Document(ten=Path(filename.replace("\\", "/")).name[:255], ma_bam=digest, duong_dan=relative, ngay_tao=now().isoformat())
    try:
        db.add(doc)
        db.flush()
        db.add_all([Chunk(tai_lieu_id=doc.id, trang=page, noi_dung=text) for page, text in chunks])
        db.commit()
    except Exception:
        db.rollback()
        path.unlink(missing_ok=True)
        raise
    return doc
